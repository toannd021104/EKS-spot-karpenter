#!/usr/bin/env python3
"""Verify presentation/karpenter-spot-finops.pptx against the deck rules (14 slides, notes, layouts, fonts, contrast, numbers, placeholders).

Usage:
  presentation/.venv/bin/python presentation/verify_deck.py [--deck PATH] [--results PATH] [--allow-synthetic]

Rules (each failure prints '[FAIL] slide N: ...'):
  * exactly 14 slides, size 12192000x6858000, layouts only Title Slide / Contents / Section slide
  * speaker notes >= 200 characters on every slide
  * no empty placeholders; every shape inside the slide and above the footer logo line
  * estimated text overflow (chars per line x lines vs box height), table height estimate
  * per-shape minimum font sizes (name-based, NFR6) and fonts Arial / Courier New only
  * text/background contrast >= 4.5 for explicitly coloured runs
  * slide 10: one native 5x4 table; banned comparative words unless all scalers are MEASURED
  * 'ước tính' on slides 9, 10, 11 and in every shape/cell that contains a dollar figure
  * cell-by-cell numbers on slides 10-12 equal the median of their own metric (n >= 1), else placeholder
  * the text 'SYNTHETIC' appears nowhere (shapes or notes) unless --allow-synthetic
Exit 0 when clean, 1 on any failure, 2 on unreadable input.
"""
import argparse
import json
import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

DECK_DIR = Path(__file__).resolve().parent
ROOT = DECK_DIR.parent
DEFAULT_DECK = DECK_DIR / "karpenter-spot-finops.pptx"
DEFAULT_RESULTS = ROOT / "results" / "summary.json"

EXPECTED_SLIDES = 14
SLIDE_W, SLIDE_H = 12192000, 6858000
ALLOWED_LAYOUTS = {"Title Slide", "Contents", "Section slide"}
ALLOWED_FONTS = {"Arial", "Courier New"}
FOOTER_TOP = int(6.62 * 914400)  # logo / footer line of the template layouts
BG_HEX = "08041B"
MIN_NOTES = 200
BANNED = ["chậm hơn", "nhanh hơn", "rẻ hơn", "đắt hơn", "tốt hơn", "ít việc nhất", "nhiều nhất", "dễ dư"]
SCALER_IDS = ["ca", "karpenter", "automode"]
TABLE_SLIDE = 10

# Minimum font size in points by shape-name prefix (NFR6); default 20.
MIN_PT_RULES = [
    ("footnote", 12), ("s02-footnote", 12), ("diagram", 16), ("s07-yaml", 14),
    ("s14-resources", 16),
]
DEFAULT_MIN_PT = 20
TABLE_MIN_PT = 14
# The line-height/glyph-width estimate runs ~10% above what LibreOffice renders for the slide-10 table
# (checked on the PNG: estimate 6.25 in vs rendered 5.7 in); calibrate so the check is not a false alarm.
TABLE_CALIBRATION = 0.9
FOOTNOTE_SUFFIX = "footnote"


def min_pt_for(name):
    if name.endswith(FOOTNOTE_SUFFIX):
        return 12
    for prefix, pt in MIN_PT_RULES:
        if name.startswith(prefix):
            return pt
    return DEFAULT_MIN_PT


# --------------------------------------------------------------------------- helpers

def lum(hex_):
    def ch(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast(a, b):
    la, lb = lum(a), lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def shape_fill_hex(shape):
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return str(shape.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def run_color_hex(run):
    try:
        if run.font.color and run.font.color.type is not None:
            return str(run.font.color.rgb)
    except (AttributeError, ValueError):
        pass
    return None


def para_text(par):
    return "".join(r.text for r in par.runs)


def frame_text(tf):
    return "\n".join(para_text(p) for p in tf.paragraphs)


def shape_text(shape):
    if shape.has_text_frame:
        return frame_text(shape.text_frame)
    if getattr(shape, "has_table", False) and shape.has_table:
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells)
    return ""


def est_frame_height_pt(tf, width_pt, default_size=20):
    """Rough text height: chars per line from average glyph width, 1.2 line spacing."""
    total = 0.0
    for par in tf.paragraphs:
        runs = par.runs
        text = para_text(par)
        size = max([r.font.size.pt for r in runs if r.font.size] or [default_size])
        mono = any((r.font.name or "") == "Courier New" for r in runs)
        factor = 0.6 if mono else 0.5
        cpl = max(1.0, width_pt / (size * factor))
        lines = max(1, math.ceil(len(text) * (1.0 if mono else 1.05) / cpl))
        total += lines * size * 1.2
    return total


def load_summary(path):
    if path is None or not Path(path).is_file():
        return None
    return json.loads(Path(path).read_text(encoding="utf-8"))


def metric(summary, scaler, group, name):
    try:
        m = summary["scalers"][scaler][group][name]
    except (KeyError, TypeError):
        return None, 0
    n = m.get("n") or 0
    if n >= 1 and m.get("median") is not None:
        return m["median"], n
    return None, 0


def get_value(summary, scaler, group, name):
    try:
        return summary["scalers"][scaler][group][name]
    except (KeyError, TypeError):
        return None


def all_measured(summary):
    if not summary:
        return False
    try:
        return all(summary["scalers"][s].get("status") == "MEASURED" for s in SCALER_IDS)
    except (KeyError, TypeError):
        return False


# --------------------------------------------------------------------------- checks

class Checker:
    def __init__(self, deck, summary, allow_synthetic):
        self.prs = Presentation(str(deck))
        self.summary = summary
        self.allow_synthetic = allow_synthetic
        self.errors = []

    def fail(self, slide_no, msg):
        where = f"slide {slide_no}" if slide_no else "deck"
        self.errors.append(f"[FAIL] {where}: {msg}")

    # -- deck level
    def check_deck(self):
        prs = self.prs
        if len(prs.slides) != EXPECTED_SLIDES:
            self.fail(0, f"expected {EXPECTED_SLIDES} slides, found {len(prs.slides)}")
        if (prs.slide_width, prs.slide_height) != (SLIDE_W, SLIDE_H):
            self.fail(0, f"slide size {prs.slide_width}x{prs.slide_height} != {SLIDE_W}x{SLIDE_H}")

    # -- per slide
    def check_slide(self, no, slide):
        layout = slide.slide_layout.name
        if layout not in ALLOWED_LAYOUTS:
            self.fail(no, f"layout '{layout}' not allowed")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if len(notes.strip()) < MIN_NOTES:
            self.fail(no, f"speaker notes {len(notes.strip())} chars < {MIN_NOTES}")
        texts = [notes]
        for shape in slide.shapes:
            self.check_shape(no, slide, shape)
            texts.append(shape_text(shape))
        full = "\n".join(texts)
        if not self.allow_synthetic and "SYNTHETIC" in full:
            self.fail(no, "text 'SYNTHETIC' found in the canonical deck")
        folded = full.casefold()
        if no in (9, 10, 11) and "ước tính" not in folded:
            self.fail(no, "literal 'ước tính' missing")
        if no == TABLE_SLIDE:
            self.check_table_slide(no, slide, texts[1:])  # shapes only: notes may name the banned words
        if no in (11, 12):
            self.check_named_numbers(no, slide)

    def check_shape(self, no, slide, shape):
        name = shape.name
        if shape.left is None or shape.top is None:
            return  # inherited geometry, nothing to measure
        left, top, w, h = shape.left, shape.top, shape.width, shape.height
        if left < 0 or top < 0 or left + w > SLIDE_W or top + h > SLIDE_H:
            self.fail(no, f"'{name}' lies outside the slide")
        layout_slide = slide.slide_layout.name in ("Title Slide", "Section slide")
        if not layout_slide and top + h > FOOTER_TOP:
            self.fail(no, f"'{name}' overlaps the footer/logo area (bottom {(top + h) / 914400:.2f} in)")
        if shape.is_placeholder and shape.has_text_frame and not frame_text(shape.text_frame).strip():
            self.fail(no, f"empty placeholder '{name}'")
        if shape.has_text_frame:
            self.check_frame(no, shape, name, shape.text_frame, w, h, shape_fill_hex(shape) or BG_HEX)
        if getattr(shape, "has_table", False) and shape.has_table:
            self.check_table_shape(no, slide, shape)

    def check_frame(self, no, shape, name, tf, w, h, bg, min_pt=None):
        text = frame_text(tf).strip()
        if text:
            width_pt = (w - tf.margin_left - tf.margin_right) / 12700
            need = est_frame_height_pt(tf, width_pt)  # insets may spill without clipping: not counted
            if need > h / 12700 * 1.1:
                self.fail(no, f"'{name}' text may overflow (needs ~{need:.0f} pt, box {h / 12700:.0f} pt)")
        floor = min_pt if min_pt is not None else min_pt_for(name)
        for par in tf.paragraphs:
            for run in par.runs:
                if not run.text.strip():
                    continue
                if run.font.size is not None and run.font.size.pt < floor:
                    self.fail(no, f"'{name}' font {run.font.size.pt:g} pt < minimum {floor} pt")
                if run.font.size is None and not shape.is_placeholder:
                    self.fail(no, f"'{name}' has a run without explicit size")
                fn = run.font.name
                if fn is not None and fn not in ALLOWED_FONTS:
                    self.fail(no, f"'{name}' uses font '{fn}'")
                col = run_color_hex(run)
                if col and contrast(col, bg) < 4.5:
                    self.fail(no, f"'{name}' contrast {contrast(col, bg):.1f} < 4.5 ({col} on {bg})")
        if "$" in text and re.search(r"\$\s*[\d_]", text) and "ước tính" not in text.casefold():
            self.fail(no, f"'{name}' has a dollar figure without 'ước tính'")

    def check_table_shape(self, no, slide, shape):
        tbl = shape.table
        col_w = [c.width for c in tbl.columns]
        total_h = 0.0
        for r_i, row in enumerate(tbl.rows):
            row_need = row.height / 12700
            for c_i, cell in enumerate(row.cells):
                bg = BG_HEX
                try:
                    if cell.fill.type == MSO_FILL.SOLID:
                        bg = str(cell.fill.fore_color.rgb)
                except (AttributeError, TypeError, ValueError):
                    pass
                tf = cell.text_frame
                self.check_frame(no, shape, f"{shape.name}[{r_i},{c_i}]", tf, col_w[c_i],
                                 10 ** 9, bg, min_pt=TABLE_MIN_PT)
                width_pt = (col_w[c_i] - cell.margin_left - cell.margin_right) / 12700
                need = est_frame_height_pt(tf, width_pt) + (cell.margin_top + cell.margin_bottom) / 12700
                row_need = max(row_need, need)
            total_h += row_need
        total_h *= TABLE_CALIBRATION
        bottom = shape.top / 12700 + total_h
        footnote_top = min([s.top for s in slide.shapes if s.name.endswith(FOOTNOTE_SUFFIX) and s.top],
                           default=FOOTER_TOP) / 12700
        if bottom > footnote_top:
            self.fail(no, f"table estimated to end at {bottom / 72:.2f} in, below the footnote at "
                          f"{footnote_top / 72:.2f} in")

    # -- slide 10
    def check_table_slide(self, no, slide, texts):
        tables = [s for s in slide.shapes if getattr(s, "has_table", False) and s.has_table]
        if len(tables) != 1:
            self.fail(no, f"expected exactly one native table, found {len(tables)}")
            return
        tbl = tables[0].table
        if (len(tbl.rows), len(tbl.columns)) != (5, 4):
            self.fail(no, f"table is {len(tbl.rows)}x{len(tbl.columns)}, expected 5x4")
            return
        full = "\n".join(texts).casefold()
        if not all_measured(self.summary):
            for word in BANNED:
                if word in full:
                    self.fail(no, f"banned comparative word '{word}' while not all scalers are MEASURED")
        for c_i, sid in enumerate(SCALER_IDS, start=1):
            self.check_cost_cell(no, tbl.cell(1, c_i).text, sid)
            self.check_seconds_cell(no, tbl.cell(2, c_i).text, sid, "scaleup", "node_provision_s",
                                    f"row 2 / {sid}")
            t3 = tbl.cell(3, c_i).text
            self.check_seconds_cell(no, t3, sid, "chaos", "recovery_s", f"row 3 / {sid}")
            self.check_pct(no, t3, sid, f"row 3 / {sid}")
        if "Đo: ___" not in "\n".join(texts) and not self.summary:
            self.fail(no, "placeholder 'Đo: ___' missing")
        if self.summary is None or all(
                metric(self.summary, s, "scaleup", "node_provision_s")[1] == 0 for s in SCALER_IDS):
            if "Ước tính (OpenCost, giá list)" not in "\n".join(texts):
                self.fail(no, "placeholder 'Ước tính (OpenCost, giá list)' missing")

    def check_cost_cell(self, no, text, sid):
        m = re.search(r"\$\s*(_+|[\d.]+)\s*/\s*giờ", text)
        if not m:
            self.fail(no, f"cost cell {sid}: no '$.../giờ' figure")
            return
        want = get_value(self.summary, sid, "cost", "opencost_per_hour_usd") if self.summary else None
        if m.group(1).startswith("_"):
            if want is not None:
                self.fail(no, f"cost cell {sid}: placeholder shown but a cost estimate exists")
        elif want is None or abs(float(m.group(1)) - want) > 0.0006:
            self.fail(no, f"cost cell {sid}: {m.group(1)} does not match summary ({want})")
        if "ước tính" not in text.casefold():
            self.fail(no, f"cost cell {sid}: missing 'ước tính'")

    def check_seconds_cell(self, no, text, sid, group, name, label):
        m = re.search(r"Đo:\s*(_+|[\d.,]+)\s*s\b", text)
        if not m:
            self.fail(no, f"{label}: no 'Đo: ... s' value")
            return
        self.compare_median(no, label, m.group(1), sid, group, name)

    def compare_median(self, no, label, token, sid, group, name):
        med, n = metric(self.summary, sid, group, name) if self.summary else (None, 0)
        if token.startswith("_"):
            if n >= 1:
                self.fail(no, f"{label}: placeholder shown but metric has n={n}")
        elif n < 1:
            self.fail(no, f"{label}: number {token} shown but metric has n=0")
        elif abs(float(token.replace(",", ".")) - med) > 0.051:
            self.fail(no, f"{label}: {token} != median {med}")

    def check_pct(self, no, text, sid, label):
        m = re.search(r"(_+|[\d.]+)\s*%\s*lỗi", text)
        if not m:
            self.fail(no, f"{label}: no '% lỗi' value")
            return
        want = get_value(self.summary, sid, "chaos", "probe_failed_pct") if self.summary else None
        if m.group(1).startswith("_"):
            if want is not None:
                self.fail(no, f"{label}: placeholder shown but failure pct exists")
        elif want is None or abs(float(m.group(1)) - want) > 0.051:
            self.fail(no, f"{label}: {m.group(1)}% does not match summary ({want})")

    # -- slides 11-12
    def check_named_numbers(self, no, slide):
        by_name = {s.name: shape_text(s) for s in slide.shapes}
        for sid in SCALER_IDS:
            if no == 11:
                self.check_named(no, by_name.get(f"s11-{sid}-prov"), sid, "scaleup", "node_provision_s", "pct0")
                self.check_named(no, by_name.get(f"s11-{sid}-all"), sid, "scaleup", "all_pods_running_s", "pct0")
            else:
                self.check_named(no, by_name.get(f"s12-{sid}-rec"), sid, "chaos", "recovery_s", "pct0")
                text = by_name.get(f"s12-{sid}-fail")
                if text is None:
                    self.fail(no, f"shape s12-{sid}-fail missing")
                else:
                    m = re.search(r"(_+|[\d.]+)\s*%", text)
                    want = get_value(self.summary, sid, "chaos", "probe_failed_pct") if self.summary else None
                    if not m:
                        self.fail(no, f"s12-{sid}-fail: no value")
                    elif m.group(1).startswith("_") and want is not None:
                        self.fail(no, f"s12-{sid}-fail: placeholder shown but value exists")
                    elif not m.group(1).startswith("_") and (want is None or abs(float(m.group(1)) - want) > 0.051):
                        self.fail(no, f"s12-{sid}-fail: {m.group(1)}% does not match summary ({want})")

    def check_named(self, no, text, sid, group, name, _unused):
        label = f"{group}.{name}/{sid}"
        if text is None:
            self.fail(no, f"shape for {label} missing")
            return
        if "___" in text:
            self.compare_median(no, label, "___", sid, group, name)
            return
        text = text.replace("SYNTHETIC", "")
        m = re.search(r"([\d.]+)\s*s\b", text)
        if not m:
            self.fail(no, f"{label}: no value in '{text}'")
            return
        self.compare_median(no, label, m.group(1), sid, group, name)


def verify(deck, results=None, allow_synthetic=False):
    """Return the list of failure strings (empty = pass)."""
    summary = load_summary(results)
    chk = Checker(deck, summary, allow_synthetic)
    chk.check_deck()
    for i, slide in enumerate(chk.prs.slides, start=1):
        chk.check_slide(i, slide)
    return chk.errors


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--deck", type=Path, default=DEFAULT_DECK)
    ap.add_argument("--results", type=Path, default=DEFAULT_RESULTS)
    ap.add_argument("--allow-synthetic", action="store_true",
                    help="accept 'SYNTHETIC' text (only for decks built under /tmp from fixtures)")
    args = ap.parse_args(argv)
    if not args.deck.is_file():
        print(f"[ERROR] deck not found: {args.deck}", file=sys.stderr)
        return 2
    try:
        errors = verify(args.deck, args.results, args.allow_synthetic)
    except (OSError, ValueError, KeyError) as exc:
        print(f"[ERROR] cannot verify: {exc}", file=sys.stderr)
        return 2
    for e in errors:
        print(e)
    if errors:
        print(f"[verify_deck] {len(errors)} failure(s)")
        return 1
    print("[verify_deck] OK: 14 slides, notes, layouts, table, fonts, contrast, numbers, placeholders")
    return 0


if __name__ == "__main__":
    sys.exit(main())
