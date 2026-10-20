#!/usr/bin/env python3
"""Tests for the synthetic-data guard and the verifier (design AC15, AC17).

Run: deck/.venv/bin/python -m unittest discover deck/tests
Every deck built here goes to a temp dir under /tmp; nothing is written to deck/ or demo/results/.
"""
import contextlib
import io
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

DECK_DIR = Path(__file__).resolve().parent.parent
ROOT = DECK_DIR.parent
sys.path.insert(0, str(DECK_DIR))

import build_deck  # noqa: E402
import verify_deck  # noqa: E402
from pptx import Presentation  # noqa: E402

SYNTHETIC = ROOT / "demo" / "scripts" / "tests" / "fixtures" / "SYNTHETIC-summary.json"
EMPTY = ROOT / "demo" / "results" / "summary.json"


def run_build(*args):
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        code = build_deck.main(list(args))
    return code, err.getvalue()


def deck_text(path):
    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                out.extend(c.text for r in shape.table.rows for c in r.cells)
        out.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


class SyntheticGuard(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="deck-test-", dir="/tmp"))
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def test_fixture_is_synthetic(self):
        import json
        self.assertTrue(json.loads(SYNTHETIC.read_text(encoding="utf-8"))["meta"]["synthetic"])

    def test_synthetic_refused_outside_tmp(self):
        out = DECK_DIR / "should-not-exist.pptx"
        code, err = run_build("--results", str(SYNTHETIC), "--out", str(out))
        self.assertNotEqual(code, 0)
        self.assertIn("synthetic", err.lower())
        self.assertFalse(out.exists())

    def test_synthetic_refused_for_results_dir(self):
        out = ROOT / "demo" / "results" / "should-not-exist.pptx"
        code, _ = run_build("--results", str(SYNTHETIC), "--out", str(out))
        self.assertNotEqual(code, 0)
        self.assertFalse(out.exists())

    def test_synthetic_under_tmp_is_prefixed_everywhere(self):
        out = self.tmp / "syn.pptx"
        code, _ = run_build("--results", str(SYNTHETIC), "--out", str(out))
        self.assertEqual(code, 0)
        text = deck_text(out)
        self.assertIn("SYNTHETIC Đo: 70 s", text)  # slide 10, Cluster Autoscaler scale-up median
        # every measured number carries the prefix: no bare 'Đo: <digit>'
        import re
        self.assertIsNone(re.search(r"(?<!SYNTHETIC )Đo:\s*\d", text))

    def test_verify_fails_on_synthetic_unless_allowed(self):
        out = self.tmp / "syn.pptx"
        self.assertEqual(run_build("--results", str(SYNTHETIC), "--out", str(out))[0], 0)
        errors = verify_deck.verify(out, SYNTHETIC, allow_synthetic=False)
        self.assertTrue(any("SYNTHETIC" in e for e in errors), errors)
        self.assertEqual(verify_deck.verify(out, SYNTHETIC, allow_synthetic=True), [])

    def test_canonical_build_has_no_synthetic_and_passes(self):
        out = self.tmp / "canon.pptx"
        self.assertEqual(run_build("--results", str(EMPTY), "--out", str(out))[0], 0)
        self.assertNotIn("SYNTHETIC", deck_text(out))
        self.assertEqual(verify_deck.verify(out, EMPTY), [])
        text = deck_text(out)
        self.assertIn("Đo: ___", text)
        self.assertIn("Ước tính (OpenCost, giá list): $___/giờ", text)


class VerifierRules(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="deck-test-", dir="/tmp"))
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.out = self.tmp / "syn.pptx"
        self.assertEqual(run_build("--results", str(SYNTHETIC), "--out", str(self.out))[0], 0)

    def tamper(self, fn):
        prs = Presentation(str(self.out))
        fn(prs)
        bad = self.tmp / "bad.pptx"
        prs.save(str(bad))
        return verify_deck.verify(bad, SYNTHETIC, allow_synthetic=True)

    def test_wrong_number_in_table_cell_is_caught(self):
        def edit(prs):
            tbl = [s for s in prs.slides[9].shapes if s.has_table][0].table
            for p in tbl.cell(2, 1).text_frame.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("70", "71")
        errors = self.tamper(edit)
        self.assertTrue(any("slide 10" in e and "median" in e for e in errors), errors)

    def test_placeholder_where_metric_has_data_is_caught(self):
        def edit(prs):
            tbl = [s for s in prs.slides[9].shapes if s.has_table][0].table
            for p in tbl.cell(2, 1).text_frame.paragraphs:
                for r in p.runs:
                    if "Đo:" in r.text:
                        r.text = "Đo: ___ s"
        errors = self.tamper(edit)
        self.assertTrue(any("placeholder shown but metric has n=" in e for e in errors), errors)

    def test_banned_word_is_caught_when_not_all_measured(self):
        def edit(prs):
            tbl = [s for s in prs.slides[9].shapes if s.has_table][0].table
            tbl.cell(4, 2).text_frame.paragraphs[0].runs[0].text = "Karpenter rẻ hơn"
        # synthetic summary: automode is TARGET, so the rule applies
        errors = self.tamper(edit)
        self.assertTrue(any("banned comparative word" in e for e in errors), errors)

    def test_dollar_without_uoc_tinh_is_caught(self):
        def edit(prs):
            shape = [s for s in prs.slides[2].shapes if s.name == "s03-analogy"][0]
            shape.text_frame.paragraphs[0].runs[0].text = "Giá $0.05 mỗi giờ"
        errors = self.tamper(edit)
        self.assertTrue(any("dollar figure" in e for e in errors), errors)

    def test_short_notes_are_caught(self):
        def edit(prs):
            prs.slides[4].notes_slide.notes_text_frame.text = "ngắn"
        errors = self.tamper(edit)
        self.assertTrue(any("slide 5" in e and "notes" in e for e in errors), errors)

    def test_small_font_is_caught(self):
        from pptx.util import Pt

        def edit(prs):
            shape = [s for s in prs.slides[3].shapes if s.name == "s04-banner"][0]
            shape.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        errors = self.tamper(edit)
        self.assertTrue(any("s04-banner" in e and "minimum" in e for e in errors), errors)


if __name__ == "__main__":
    unittest.main()
